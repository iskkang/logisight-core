#!/usr/bin/env python3
"""Download, validate and ingest the 2026 EU Combined Nomenclature.

CN 2026 is legally established by Commission Implementing Regulation
(EU) 2025/1926. Download mirrors are used only as machine-readable
distributions; every row keeps the EU legal source as its provenance.

Required env:
  SUPABASE_URL
  SUPABASE_SERVICE_ROLE_KEY

Optional env:
  EU_CN_2026_XLSX_URL
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
from datetime import date
from html.parser import HTMLParser
from typing import Iterable
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook

LEGAL_SOURCE_URL = "https://eur-lex.europa.eu/eli/reg_impl/2025/1926/oj"
SOURCE_VERSION = "CN 2026 / Commission Implementing Regulation (EU) 2025/1926"
CODE_RE = re.compile(r"^\d{8}$")
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; Logisight-CN-Ingest/1.0; +https://logisight.net)"
}

# Malta's NSO currently returns 403 to GitHub-hosted runners. Keep it as a
# fallback, but discover an XLSX first from Spain's official Tax Agency page.
LANDING_PAGES = [
    "https://sede.agenciatributaria.gob.es/Sede/en_gb/estadisticas/estadisticas-comercio-exterior/nomenclatura-combinada-ano.html",
]
DIRECT_FALLBACKS = [
    "https://nso.gov.mt/wp-content/uploads/CN_Validity_2026-483e17c9fb3ffdee.xlsx",
]


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[tuple[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "a":
            self._href = dict(attrs).get("href")
            self._text = []

    def handle_data(self, data: str) -> None:
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a" and self._href is not None:
            self.links.append((self._href, " ".join(self._text).strip()))
            self._href = None
            self._text = []


def discover_xlsx_urls() -> list[str]:
    explicit = os.getenv("EU_CN_2026_XLSX_URL")
    if explicit:
        return [explicit]

    discovered: list[str] = []
    for page in LANDING_PAGES:
        try:
            response = requests.get(page, headers=HEADERS, timeout=30)
            response.raise_for_status()
            parser = LinkParser()
            parser.feed(response.text)
            for href, label in parser.links:
                haystack = f"{href} {label}".lower()
                if (
                    (".xlsx" in haystack or ".xls" in haystack)
                    and ("2026" in haystack)
                    and ("nomenclature" in haystack or "estructura" in haystack or "structure" in haystack)
                    and "correspond" not in haystack
                    and "review" not in haystack
                ):
                    discovered.append(urljoin(page, href))
        except requests.RequestException as exc:
            print(f"Source discovery failed for {page}: {exc}", file=sys.stderr)

    return list(dict.fromkeys(discovered + DIRECT_FALLBACKS))


def download_workbook(path: str) -> str:
    errors: list[str] = []
    for url in discover_xlsx_urls():
        try:
            response = requests.get(url, headers=HEADERS, timeout=60)
            response.raise_for_status()
            content_type = response.headers.get("content-type", "").lower()
            if len(response.content) < 10_000:
                raise RuntimeError("response is unexpectedly small")
            if response.content[:2] != b"PK" and "spreadsheet" not in content_type and "excel" not in content_type:
                raise RuntimeError(f"response is not an XLSX workbook ({content_type})")
            with open(path, "wb") as handle:
                handle.write(response.content)
            print(f"Downloaded CN 2026 workbook from {url}")
            return url
        except (requests.RequestException, RuntimeError) as exc:
            errors.append(f"{url}: {exc}")
    raise RuntimeError("All CN 2026 download sources failed:\n" + "\n".join(errors))


def normalize_code(value: object) -> str:
    """Return an 8-digit CN code only when the source itself contains 8 digits.

    Never zero-pad shorter HS hierarchy codes: doing so turns heading 0302 into
    the fake commodity code 00000302.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = str(int(value))
    elif isinstance(value, int):
        value = str(value)
    digits = re.sub(r"\D", "", str(value))
    return digits if len(digits) == 8 else ""


def find_columns(rows: list[tuple[object, ...]]) -> tuple[int, int, int]:
    code_names = {"cn code", "cn_code", "code", "commodity code", "cn8", "código nc", "codigo nc"}
    desc_names = {"description", "goods description", "cn description", "text", "descripción", "descripcion"}

    for row_idx, row in enumerate(rows[:30]):
        normalized = [str(v or "").strip().lower() for v in row]
        code_col = next((i for i, v in enumerate(normalized) if v in code_names), None)
        desc_col = next((i for i, v in enumerate(normalized) if v in desc_names), None)
        if code_col is not None and desc_col is not None:
            return row_idx, code_col, desc_col

    sample = rows[:250]
    width = max((len(r) for r in sample), default=0)
    for code_col in range(width):
        hits = sum(1 for r in sample if code_col < len(r) and CODE_RE.fullmatch(normalize_code(r[code_col])))
        if hits < 20:
            continue
        for desc_col in range(width):
            if desc_col == code_col:
                continue
            values = [str(r[desc_col] or "").strip() for r in sample if desc_col < len(r) and str(r[desc_col] or "").strip()]
            text_hits = sum(1 for value in values if len(value) >= 5 and re.search(r"[A-Za-zÀ-ÿ]", value) and not re.fullmatch(r"[\\d\\s./_-]+", value))
            if text_hits >= max(20, hits // 2):
                return 0, code_col, desc_col
    raise RuntimeError("Could not identify CN code/description columns; refusing to ingest")


def parse_workbook(path: str) -> list[dict[str, object]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    candidates: list[dict[str, object]] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        try:
            header_row, code_col, desc_col = find_columns(rows)
        except RuntimeError:
            continue
        for row in rows[header_row + 1:]:
            if max(code_col, desc_col) >= len(row):
                continue
            code = normalize_code(row[code_col])
            description = str(row[desc_col] or "").strip()
            if not CODE_RE.fullmatch(code) or len(description) < 2:
                continue
            if not re.search(r"[A-Za-zÀ-ÿ]", description) or re.fullmatch(r"[\\d\\s./_-]+", description):
                continue
            candidates.append({
                "market": "EU", "nomenclature": "CN", "code": code,
                "parent_code": code[:6], "description": description, "level": 8,
                "valid_from": date(2026, 1, 1).isoformat(),
                "valid_to": date(2026, 12, 31).isoformat(),
                "source_name": "EU Combined Nomenclature 2026",
                "source_url": LEGAL_SOURCE_URL, "source_version": SOURCE_VERSION,
                "is_active": True, "is_leaf": True,
            })

    rows = list({row["code"]: row for row in candidates}.values())
    # Statistics Estonia publishes 9,791 level-5 commodity items for CN 2026.
    # Keep a conservative lower bound so layout/source changes fail closed.
    if len(rows) < 9_000:
        raise RuntimeError(f"Parsed only {len(rows)} unique CN8 rows; refusing suspicious dataset")
    if any(not CODE_RE.fullmatch(str(row["code"])) for row in rows):
        raise RuntimeError("Invalid CN8 code detected")
    textual = sum(1 for row in rows if re.search(r"[A-Za-zÀ-ÿ]", str(row["description"])) and not re.fullmatch(r"[\\d\\s./_-]+", str(row["description"])))
    if textual < int(len(rows) * 0.98):
        raise RuntimeError(f"Only {textual}/{len(rows)} rows have textual descriptions; refusing to ingest")
    return rows


def chunks(items: list[dict[str, object]], size: int = 500) -> Iterable[list[dict[str, object]]]:
    for i in range(0, len(items), size):
        yield items[i:i + size]


def upsert(rows: list[dict[str, object]]) -> None:
    base = os.environ["SUPABASE_URL"].rstrip("/")
    key = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    url = f"{base}/rest/v1/customs_nomenclature?on_conflict=market,nomenclature,code,valid_from"
    headers = {
        "apikey": key, "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates,return=minimal",
    }
    for batch in chunks(rows):
        response = requests.post(url, headers=headers, json=batch, timeout=60)
        response.raise_for_status()


def main() -> int:
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
        download_workbook(tmp.name)
        rows = parse_workbook(tmp.name)
    upsert(rows)
    print(f"Ingested {len(rows)} validated CN8 rows for 2026")
    return 0


if __name__ == "__main__":
    sys.exit(main())
