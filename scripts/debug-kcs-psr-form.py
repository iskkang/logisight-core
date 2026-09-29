import io,requests,zipfile
from openpyxl import load_workbook
url="https://www.customs.go.kr/common/nttFileDownload.do?fileKey=baee47db73be787e49c4e253fc2f5c23"
r=requests.get(url,headers={"User-Agent":"Mozilla/5.0"},timeout=120); r.raise_for_status()
print("STATUS",r.status_code,"BYTES",len(r.content),"CT",r.headers.get("content-type"),"CD",r.headers.get("content-disposition"))
print("MAGIC",r.content[:16])
try:
    wb=load_workbook(io.BytesIO(r.content),read_only=True,data_only=True)
    print("SHEETS",wb.sheetnames)
    for ws in wb.worksheets:
        print("SHEET",ws.title,"MAX",ws.max_row,ws.max_column)
        for i,row in enumerate(ws.iter_rows(values_only=True),1):
            print("ROW",i,row[:12])
            if i>=12:break
except Exception as e:
    print("XLSX ERROR",repr(e))
