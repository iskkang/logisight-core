#!/usr/bin/env python3
import io, os, re, requests
from collections import defaultdict
from datetime import date, datetime
from openpyxl import load_workbook
import xml.etree.ElementTree as ET

BASE="https://circabc.europa.eu/service/api/node/workspace/SpacesStore"
ROOT="64db9d0f-e7c9-4084-afe9-f47e70e53c10"
LIBRARY="https://circabc.europa.eu/ui/group/0e5f18c2-4b2f-42e9-aed4-dfe50ae1263b/library/64db9d0f-e7c9-4084-afe9-f47e70e53c10"
HEAD={"Authorization":"Basic Z3Vlc3Q6","User-Agent":"Logisight-EU-TARIC-FTA/1.0"}
MEASURE_TYPE="142"
ORIGIN="KR"

def children(node):
    r=requests.get(f"{BASE}/{node}/children",headers=HEAD,timeout=60); r.raise_for_status()
    root=ET.fromstring(r.content)
    ns={"a":"http://www.w3.org/2005/Atom"}
    out=[]
    for e in root.findall("a:entry",ns):
        title=(e.findtext("a:title",default="",namespaces=ns) or "").strip()
        content=e.find("a:content",ns); mime=content.attrib.get("type","") if content is not None else ""
        nid=None
        for link in e.findall("a:link",ns):
            if link.attrib.get("rel")=="self":
                m=re.search(r"SpacesStore/i/([0-9a-f-]{36})",link.attrib.get("href",""))
                if m: nid=m.group(1)
        if nid: out.append({"title":title,"id":nid,"mime":mime})
    return out

def latest_month():
    years=[x for x in children(ROOT) if re.fullmatch(r"20\d\d",x["title"]) and not x["mime"]]
    year=max(years,key=lambda x:x["title"])
    months=[x for x in children(year["id"]) if re.match(r"^\d{2}\s*-",x["title"]) and not x["mime"]]
    month=max(months,key=lambda x:x["title"][:2])
    return year["title"],month

def download(node):
    r=requests.get(f"{BASE}/{node}/content",headers=HEAD,timeout=180); r.raise_for_status(); return r.content

def find_month_file(month_id, needle):
    needle=needle.lower()
    for x in children(month_id):
        if needle in x["title"].lower() and ("spreadsheet" in x["mime"] or "excel" in x["mime"]):
            return x
    return None

def korea_origin_codes(month_id, today):
    geo=find_month_file(month_id,"geographical area composition")
    codes={"KR"}
    if not geo: return codes
    wb=load_workbook(io.BytesIO(download(geo["id"])),read_only=True,data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if len(row)<10: continue
            member_iso=str(row[6] or "").strip().upper()
            if member_iso!="KR": continue
            start=as_date(row[8]); end=as_date(row[9])
            if start and start>today: continue
            if end and end<today: continue
            group=str(row[0] or "").strip()
            if group: codes.add(group)
    return codes

def norm_header(v):
    return re.sub(r"[^a-z0-9]+"," ",str(v or "").strip().lower()).strip()

def find_header(ws):
    aliases={
      "code":["goods code","goods nomenclature code","goods nomenclature item id"],
      "start":["validity start date","start date"],
      "end":["validity end date","end date"],
      "origin":["origin code","geographical area code"],
      "type":["measure type code"],
      "duty":["duty"],
      "legal":["legal reference","legal act"],
    }
    for rn,row in enumerate(ws.iter_rows(values_only=True),1):
        vals=[norm_header(x) for x in row]
        if rn>30: break
        idx={}
        for k,names in aliases.items():
            for n in names:
                if n in vals: idx[k]=vals.index(n); break
        if all(k in idx for k in ("code","origin","type","duty")): return rn,idx
    # Official TARIC extraction fixed column fallback from DG TAXUD documentation:
    # A goods code, D start, E end, G origin, I legal ref, J duty, K origin code, L measure type code.
    return 0,{"code":0,"start":3,"end":4,"origin":10,"type":11,"legal":8,"duty":9}

def as_date(v):
    if v is None or v=="": return None
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    s=str(v).strip()
    for f in ("%Y-%m-%d","%d/%m/%Y","%d.%m.%Y"):
        try:return datetime.strptime(s,f).date()
        except: pass
    return None

PCT=re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*%?\s*$")
def parse_percent(duty):
    s=str(duty or "").strip()
    m=PCT.match(s)
    return float(m.group(1).replace(",",".")) if m else None

def fetch_cn8(base,key):
    out=[]; offset=0
    h={"apikey":key,"Authorization":f"Bearer {key}"}
    while True:
        q=(f"{base}/rest/v1/customs_nomenclature?select=code"
           f"&market=eq.EU&nomenclature=eq.CN&is_active=eq.true&is_leaf=eq.true"
           f"&order=code.asc&offset={offset}&limit=1000")
        r=requests.get(q,headers=h,timeout=60); r.raise_for_status(); batch=r.json()
        out += [re.sub(r"\D","",x["code"])[:8] for x in batch if x.get("code")]
        if len(batch)<1000: break
        offset += 1000
    return sorted(set(x for x in out if len(x)==8))

def main():
    sb=os.environ["SUPABASE_URL"].rstrip("/"); key=os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    year,month=latest_month()
    files=[x for x in children(month["id"]) if "duties import" in x["title"].lower() and ("spreadsheet" in x["mime"] or "excel" in x["mime"])]
    if not files: raise RuntimeError(f"No Duties Import XLSX in {year}/{month['title']}")
    today=date.today()
    origin_codes=korea_origin_codes(month["id"],today)
    print("KR applicable TARIC geographical-area codes:", sorted(origin_codes))
    cn8=fetch_cn8(sb,key)
    if len(cn8)<8000: raise RuntimeError(f"Only {len(cn8)} active CN8 codes; refusing import")

    raw=[]
    for f in files:
        wb=load_workbook(io.BytesIO(download(f["id"])),read_only=True,data_only=True)
        for ws in wb.worksheets:
            hr,idx=find_header(ws)
            for row in ws.iter_rows(min_row=(hr+1 if hr else 1),values_only=True):
                if len(row)<=max(idx.values()): continue
                typ=str(row[idx["type"]] or "").strip()
                origin=str(row[idx["origin"]] or "").strip().upper()
                if typ!=MEASURE_TYPE or origin not in origin_codes: continue
                code=re.sub(r"\D","",str(row[idx["code"]] or ""))
                if len(code)<2: continue
                start=as_date(row[idx.get("start",3)]); end=as_date(row[idx.get("end",4)])
                if start and start>today: continue
                if end and end<today: continue
                duty=str(row[idx["duty"]] or "").strip()
                if not duty: continue
                raw.append({"source_code":code,"duty":duty,"rate":parse_percent(duty),
                            "legal":str(row[idx.get("legal",8)] or "").strip(),
                            "start":start.isoformat() if start else None,"end":end.isoformat() if end else None,
                            "file":f["title"]})
    if len(raw)<500: raise RuntimeError(f"Only {len(raw)} active KR-applicable tariff-preference measures parsed; refusing mutation")

    by_cn=defaultdict(list)
    for m in raw:
        code=m["source_code"]
        prefix=code[:8] if len(code)>=8 else code
        targets=[c for c in cn8 if c.startswith(prefix)]
        for c in targets: by_cn[c].append(m)

    payload=[]
    for c,items in by_cn.items():
        depth=max(min(len(x["source_code"]),10) for x in items)
        best=[x for x in items if min(len(x["source_code"]),10)==depth]
        duties=sorted(set(x["duty"] for x in best))
        rates=sorted(set(x["rate"] for x in best if x["rate"] is not None))
        uniform=len(duties)==1
        rate=rates[0] if uniform and len(rates)==1 else None
        rate_text=duties[0] if uniform else "Multiple TARIC subline rates: "+" | ".join(duties[:8])
        legal="; ".join(sorted(set(x["legal"] for x in best if x["legal"])))[:1000] or "EU-Korea FTA / TARIC measure 142"
        payload.append({
          "cn_code":c,"origin_country":"KR","destination_country":None,"measure_type":"PREFERENCE",
          "rate_percent":rate,"rate_text":rate_text,"title":"EU-Korea FTA tariff preference",
          "detail":"Official TARIC measure type 142 for Korean origin. Preferential treatment applies only when the EU-Korea FTA origin requirements and proof/declaration requirements are satisfied.",
          "legal_basis":legal,"source_url":LIBRARY,
          "valid_from":min((x["start"] for x in best if x["start"]),default=f"{year}-01-01"),
          "valid_to":max((x["end"] for x in best if x["end"]),default=f"{year}-12-31"),
          "is_active":True,
          "metadata":{"source":"EU TARIC KR preference","year":int(year),"measure_type":"142","source_month":month["title"],
                      "source_codes":sorted(set(x["source_code"] for x in best))[:20],"mixed_taric_subline_rates":not uniform}
        })
    if len(payload)<7000: raise RuntimeError(f"Only {len(payload)} CN8 KR preference rows resolved; refusing mutation")

    h={"apikey":key,"Authorization":f"Bearer {key}","Content-Type":"application/json","Prefer":"return=minimal"}
    # Replace KR preference snapshot only after the new snapshot has passed all guards.
    q=f"{sb}/rest/v1/eu_customs_measures?measure_type=eq.PREFERENCE&origin_country=eq.KR"
    requests.delete(q,headers=h,timeout=60).raise_for_status()
    for i in range(0,len(payload),400):
        requests.post(f"{sb}/rest/v1/eu_customs_measures",headers=h,json=payload[i:i+400],timeout=60).raise_for_status()
    zero=sum(1 for x in payload if x["rate_percent"]==0)
    mixed=sum(1 for x in payload if x["rate_percent"] is None)
    print(f"Ingested {len(payload)} KR→EU CN8 preference rows from official TARIC {year}/{month['title']}; 0%={zero}; complex/mixed={mixed}")

if __name__=="__main__": main()
