#!/usr/bin/env python3
"""Ingest 2026 EU CN8 conventional (third-country) duties from Commission TARIC raw Excel.

Set TARIC_RAW_XLSX_URL to the current official Commission/CIRCABC Excel download.
The parser fails closed if it cannot identify >=9,000 CN8 rows or the duty column.
"""
import io, os, re, requests
from datetime import date
from openpyxl import load_workbook

SOURCE_PAGE="https://taxation-customs.ec.europa.eu/online-services/online-services-and-databases-customs/eu-customs-tariff-taric_en"
CODE=re.compile(r"^\d{8}$")
PCT=re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*%\s*$")

def norm(v):
    d=re.sub(r"\D","",str(v or ""))
    return d[:8] if len(d)>=8 else ""

def find_cols(rows):
    for i,row in enumerate(rows[:50]):
        vals=[str(x or "").strip().lower() for x in row]
        ci=next((j for j,v in enumerate(vals) if v in {"cn code","commodity code","goods code","code nc","código nc"}),None)
        di=next((j for j,v in enumerate(vals) if "third" in v and ("duty" in v or "country" in v) or v in {"conventional rate of duty","duty rate"}),None)
        if ci is not None and di is not None:return i,ci,di
    raise RuntimeError("Could not identify CN8 / third-country duty columns")

def parse(blob):
    wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    out={}
    for ws in wb.worksheets:
        rows=list(ws.iter_rows(values_only=True))
        try:h,ci,di=find_cols(rows)
        except RuntimeError:continue
        for row in rows[h+1:]:
            if max(ci,di)>=len(row):continue
            code=norm(row[ci]); rate_text=str(row[di] or "").strip()
            if not CODE.fullmatch(code) or not rate_text:continue
            free=rate_text.lower() in {"free","0","0%","0 %"}
            m=PCT.fullmatch(rate_text)
            rate=0.0 if free else (float(m.group(1).replace(",",".")) if m else None)
            out[code]={"cn_code":code,"origin_country":None,"destination_country":None,"measure_type":"THIRD_COUNTRY_DUTY","rate_percent":rate,"rate_text":rate_text,"title":"EU CN 2026 conventional rate of duty","detail":"Official TARIC/CCT third-country duty; complex rates are preserved verbatim in rate_text.","legal_basis":"Council Regulation (EEC) No 2658/87; CN 2026","source_url":SOURCE_PAGE,"valid_from":date(2026,1,1).isoformat(),"valid_to":date(2026,12,31).isoformat(),"is_active":True,"metadata":{"source":"EU TARIC raw data","year":2026}}
    if len(out)<9000:raise RuntimeError(f"Only {len(out)} CN8 duty rows parsed; refusing mutation")
    return list(out.values())

def main():
    url=os.environ["TARIC_RAW_XLSX_URL"]
    r=requests.get(url,timeout=120,headers={"User-Agent":"Logisight-TARIC-Ingest/1.0"});r.raise_for_status()
    rows=parse(r.content)
    base=os.environ["SUPABASE_URL"].rstrip("/"); key=os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    headers={"apikey":key,"Authorization":f"Bearer {key}","Content-Type":"application/json","Prefer":"return=minimal"}
    # Replace only this importer's 2026 conventional-duty snapshot.
    q=f'{base}/rest/v1/eu_customs_measures?measure_type=eq.THIRD_COUNTRY_DUTY&valid_from=eq.2026-01-01&metadata->>source=eq.EU%20TARIC%20raw%20data'
    requests.delete(q,headers=headers,timeout=60).raise_for_status()
    for i in range(0,len(rows),500):
        requests.post(f"{base}/rest/v1/eu_customs_measures",headers=headers,json=rows[i:i+500],timeout=60).raise_for_status()
    print(f"Ingested {len(rows)} EU CN8 conventional-duty rows")
if __name__=="__main__":main()
